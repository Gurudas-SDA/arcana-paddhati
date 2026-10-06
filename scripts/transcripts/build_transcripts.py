"""Bundle the lecture transcripts the book links to as offline pages.

WHY: the book's sources («транскрипт») pointed to Google Drive. Offline (and on
Android, where the Drive app intercepts the link) they did not open. Satkirti's
rule of 06.10.2026: every transcript the book refers to opens OFFLINE inside the
app (audio stays external).

WHAT: for every `transcript_url` in data/book*.json (Google Drive docx) this
  1. downloads the docx (Drive export; cached in --cache),
  2. converts it to a clean, book-styled static HTML page
     public/transcripts/<nr>-<lang>.html (paragraphs, bold/italic, verse
     tables sanskrit | translation, separators),
  3. writes lib/transcripts.json  {drive file id: page name}
which components/MoodBlock.tsx / SectionContent.tsx use to point the
«транскрипт» link to the internal page. public/ is precached by the service
worker (scripts/build-precache.mjs), so the pages work offline.

HOW: python scripts/transcripts/build_transcripts.py [--cache DIR] [--refresh]
Run after new sources are added to the book, then commit public/transcripts/
and lib/transcripts.json.
"""
import argparse
import glob
import html
import json
import os
import re
import subprocess
import sys

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, "public", "transcripts")
MAP = os.path.join(ROOT, "lib", "transcripts.json")
BASE = "/arcana-paddhati/"

ID_RE = re.compile(r"/d/([\w-]+)|[?&]id=([\w-]+)")


def drive_id(url):
    m = ID_RE.search(url or "")
    return (m.group(1) or m.group(2)) if m else None


def collect():
    """{drive id: {"nrs": set, "titles": set, "dates": set}} from the books."""
    found = {}

    def walk(o):
        if isinstance(o, dict):
            url = o.get("transcript_url")
            if isinstance(url, str):
                i = drive_id(url)
                if i:
                    e = found.setdefault(i, {"nrs": set(), "titles": set(), "dates": set()})
                    if o.get("nr"):
                        e["nrs"].add(str(o["nr"]))
                    if o.get("title"):
                        e["titles"].add(o["title"])
                    if o.get("date"):
                        e["dates"].add(o["date"])
            for v in o.values():
                walk(v)
        elif isinstance(o, list):
            for v in o:
                walk(v)

    for f in sorted(glob.glob(os.path.join(ROOT, "data", "book*.json"))):
        with open(f, encoding="utf8") as fh:
            walk(json.load(fh))
    return found


def download(i, cache, refresh):
    p = os.path.join(cache, f"{i}.docx")
    if refresh or not (os.path.exists(p) and os.path.getsize(p) > 1000):
        subprocess.run(["curl", "--ssl-no-revoke", "-sL", "-o", p,
                        f"https://drive.google.com/uc?export=download&id={i}"], check=True)
    with open(p, "rb") as fh:
        if fh.read(2) != b"PK":
            raise RuntimeError(f"{i}: not a docx (Drive refused?)")
    return p


def detect_lang(text):
    cyr = len(re.findall(r"[А-Яа-яЁё]", text))
    lat = len(re.findall(r"[A-Za-z]", text))
    if cyr > lat:
        return "uk" if len(re.findall(r"[ієїґІЄЇҐ]", text)) > cyr * 0.01 else "ru"
    return "en"


def runs_html(p):
    out = []
    for r in p.runs:
        t = html.escape(r.text).replace("\n", "<br>")
        if not t:
            continue
        if r.bold:
            t = f"<b>{t}</b>"
        if r.italic:
            t = f"<i>{t}</i>"
        out.append(t)
    return "".join(out)


SEP = re.compile(r"^[─━—\-_=]{6,}\s*(\[\d+\])?$")
SEP_LABEL = re.compile(r"^[─━]{2,}\s*(.+?)\s*[─━]{2,}$")
IAST = re.compile(r"[āīūṛṝḷṅñṭḍṇśṣṁṃḥ]")


def cell_html(cell, sanskrit):
    lines = [runs_html(p) for p in cell.paragraphs]
    body = "<br>".join(x for x in lines if x.strip())
    lang = ' lang="sa-Latn"' if sanskrit else ""
    return f"<td{lang}>{body}</td>"


def convert(path, title, lang):
    d = docx.Document(path)
    parts = []
    para_buf = []  # consecutive short lines -> one verse block

    def flush():
        if para_buf:
            sa = all(IAST.search(x[1]) or len(x[1]) < 60 for x in para_buf) and len(para_buf) >= 2
            if sa:
                parts.append('<p class="verse" lang="sa-Latn">' + "<br>".join(x[0] for x in para_buf) + "</p>")
            else:
                parts.extend(f"<p>{x[0]}</p>" for x in para_buf)
            para_buf.clear()

    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, d)
            text = p.text.strip()
            if not text:
                flush()
                continue
            m = SEP_LABEL.match(text)
            if m and not SEP.match(text):
                flush()
                parts.append(f'<p class="sep">{html.escape(m.group(1))}</p>')
                continue
            if SEP.match(text):
                flush()
                parts.append("<hr>")
                continue
            h = runs_html(p)
            if len(text) < 70 and IAST.search(text):
                para_buf.append((h, text))
            else:
                flush()
                parts.append(f"<p>{h}</p>")
        elif tag == "tbl":
            flush()
            t = Table(el, d)
            rows = []
            for r in t.rows:
                cells = r.cells
                rows.append("<tr>" + "".join(cell_html(c, k == 0 and len(cells) > 1) for k, c in enumerate(cells)) + "</tr>")
            parts.append('<table class="verses">' + "".join(rows) + "</table>")
    flush()
    back = {"ru": "Назад к книге", "uk": "Назад до книги"}.get(lang, "Back to the book")
    note = {"ru": "Транскрипт лекции (сохранён в приложении, открывается без интернета)",
            "uk": "Транскрипт лекції (збережено в застосунку, відкривається без інтернету)"}.get(
        lang, "Lecture transcript (stored in the app, opens offline)")
    return f"""<!doctype html>
<html lang="{lang}" translate="no">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<meta name="google" content="notranslate">
<title>{html.escape(title)}</title>
<style>
:root {{ color-scheme: light; }}
body {{ margin: 0; background: #FDF8F0; color: #2C1810; font-family: 'Noto Serif', Georgia, 'Times New Roman', serif;
  font-size: clamp(17px, calc(15.5px + 0.5vw), 21px); line-height: 1.6; -webkit-text-size-adjust: 100%; }}
header {{ position: sticky; top: 0; background: #FDF8F0; border-bottom: 1px solid #E8DCC8; padding: max(8px, env(safe-area-inset-top)) 12px 8px; display: flex; align-items: center; gap: 8px; z-index: 1; }}
header a {{ display: inline-flex; align-items: center; min-height: 44px; min-width: 44px; padding: 0 10px; color: #8B6508; text-decoration: none; font-size: 16px; border-radius: 10px; }}
header a:active {{ background: #F5E6C8; }}
main {{ max-width: 46em; margin: 0 auto; padding: 16px max(16px, 4vw) calc(48px + env(safe-area-inset-bottom)); }}
h1 {{ font-size: 1.25em; line-height: 1.3; margin: 0.4em 0 0.2em; }}
.note {{ color: #8B6508; font-size: 0.8em; margin: 0 0 1.2em; }}
p {{ margin: 0 0 0.8em; }}
.verse, td[lang] {{ font-weight: 600; -webkit-touch-callout: none; }}
.verse {{ text-align: center; margin: 1em 0; }}
.sep {{ text-align: center; color: #8B6508; font-size: 0.75em; letter-spacing: 0.08em; margin: 1.4em 0 0.4em; }}
hr {{ border: 0; border-top: 1px solid #E8DCC8; margin: 1.2em 0; }}
table.verses {{ width: 100%; border-collapse: collapse; margin: 1em 0; font-size: 0.92em; }}
table.verses td {{ vertical-align: top; padding: 8px 10px; border-top: 1px solid #E8DCC8; }}
@media (max-width: 639px) {{
  table.verses, table.verses tbody, table.verses tr, table.verses td {{ display: block; width: auto; }}
  table.verses td + td {{ border-top: 0; padding-top: 0; color: #5C3D2E; }}
}}
</style>
</head>
<body translate="no">
<header><a href="{BASE}" onclick="if (history.length > 1) {{ history.back(); return false; }}">‹ {back}</a></header>
<main>
<h1>{html.escape(title)}</h1>
<p class="note">{note}</p>
{chr(10).join(parts)}
</main>
</body>
</html>
"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cache", default=os.path.join(ROOT, "scripts", "transcripts", "cache"))
    ap.add_argument("--refresh", action="store_true")
    a = ap.parse_args()
    os.makedirs(a.cache, exist_ok=True)
    os.makedirs(OUT, exist_ok=True)
    found = collect()
    names, mapping, total = set(), {}, 0
    for i in sorted(found):
        e = found[i]
        path = download(i, a.cache, a.refresh)
        d = docx.Document(path)
        sample = " ".join(p.text for p in d.paragraphs[:200])
        lang = detect_lang(sample)
        nr = sorted(e["nrs"])[0] if e["nrs"] else i[:10]
        name = f"{nr}-{lang}"
        k = 2
        while name in names:
            name = f"{nr}-{lang}-{k}"
            k += 1
        names.add(name)
        title = sorted(e["titles"])[0] if e["titles"] else f"№ {nr}"
        if e["nrs"]:
            title = f"{title} · № {nr}"
        page = convert(path, title, lang)
        with open(os.path.join(OUT, f"{name}.html"), "w", encoding="utf8", newline="\n") as fh:
            fh.write(page)
        total += len(page.encode("utf8"))
        mapping[i] = name
    for f in os.listdir(OUT):
        if f.endswith(".html") and f[:-5] not in names:
            os.remove(os.path.join(OUT, f))
    with open(MAP, "w", encoding="utf8", newline="\n") as fh:
        json.dump(mapping, fh, ensure_ascii=False, indent=0, sort_keys=True)
        fh.write("\n")
    print(f"transcripts: {len(mapping)} pages, {total / 1e6:.2f} MB -> public/transcripts/, lib/transcripts.json")


if __name__ == "__main__":
    sys.exit(main())
